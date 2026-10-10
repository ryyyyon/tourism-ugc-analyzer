"""
分類の精度を検証するための、ラベル付け用Excelを作る

使い方:
  python make_validation_sample.py                  # 各地域から国内15本・海外15本ずつ（計90本）
  python make_validation_sample.py --per-group 10   # 本数を減らす
  python make_validation_sample.py --themes グルメ 祭り・イベント 歴史・寺社 --per-theme 15
      # テーマ別の精度を確かめる：海外向けの動画のうち、分類でそのテーマとした動画を15本ずつ選ぶ
      # （企画書の「食と祭りは寺の約4倍」に使うテーマ。前の検証で選んだ動画は除く）

作られるファイル（data/validation/）:
  labeling_<日時>.xlsx     … 人が正解を記入するファイル。分類の結果は載せていない
  predictions_<日時>.csv   … 分類の結果。記入が終わるまで開かない

【設計の考え方】
  - 分類の結果を見せずに正解をつけてもらう（先に結果が見えると「合っている」と判断しやすくなるため）
  - 銚子は海外向けの動画が少ないので、地域ごとに国内向け・海外向けを同じ本数ずつ選ぶ
    （精度は国内・海外それぞれで出す。全体の精度を出すときは、この偏りに注意）
  - ノイズとして除外した動画も含めて選ぶ（除外の判断が正しいかも検証するため）
  - 乱数のシードを固定し、同じ条件なら同じ動画が選ばれるようにする
  - 並び順はシャッフルする（地域ごとに固まっていると、判断の基準がずれやすいため）

記入後は evaluate_validation.py で精度を計算する（テーマを指定して選んだときは --all を付けて、前の検証とまとめる）。
"""

import argparse
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from classify_themes import build_themed_videos
from region_config import DATA_ROOT, list_regions
from theme_rules import TOURISM_THEMES

def _all_theme_choices():
    """共通のテーマに、地域の設定ファイルで追加したテーマ（空港、乗り継ぎ・待ち時間など）を足す"""
    from region_config import load_region
    themes = list(TOURISM_THEMES.keys())
    for slug in list_regions():
        for t in load_region(slug).get("extra_tourism", {}):
            if t not in themes:
                themes.append(t)
    return themes + ["旅行全般", "どれにも当てはまらない"]


THEME_CHOICES = _all_theme_choices()
TOURISM_CHOICES = ["観光", "観光以外"]
LANG_CHOICES = ["日本語", "英語", "韓国語", "中国語", "タイ語", "スペイン語", "その他"]

FONT = "Arial"
YELLOW = PatternFill("solid", start_color="FFFF00")
HEADER_FILL = PatternFill("solid", start_color="D9D9D9")


def already_sampled(out_dir):
    """これまでの検証で選んだ動画（同じ動画を二度記入しないように除く）"""
    ids = set()
    for p in out_dir.glob("predictions_*.csv"):
        ids.update(pd.read_csv(p)["video_id"].astype(str))
    return ids


def sample_by_theme(themes, per_theme, market, seed, regions, exclude):
    """分類でそのテーマとした動画を、テーマごとに選ぶ（テーマ別の適合率を確かめる用）"""
    frames = []
    for slug in regions:
        region, videos, _, _ = build_themed_videos(slug)
        frames.append(videos.assign(region=region["name"].split("（")[0]))
    allv = pd.concat(frames, ignore_index=True)
    allv = allv[~allv["is_noise"] & ~allv["video_id"].astype(str).isin(exclude)]
    if market != "すべて":
        allv = allv[allv["market"] == market]
    picks = []
    for th in themes:
        d = allv[allv["primary_theme"] == th]
        if d.empty:
            print(f"（{th}に分類された動画がありません）")
            continue
        picks.append(d.sample(n=min(per_theme, len(d)), random_state=seed))
    df = pd.concat(picks).sample(frac=1, random_state=seed).reset_index(drop=True)
    df["番号"] = range(1, len(df) + 1)
    return df


def sample_videos(per_group, seed, regions):
    picks = []
    for slug in regions:
        region, videos, _, _ = build_themed_videos(slug)
        name = region["name"].split("（")[0]
        videos = videos.assign(region=name)
        dom = videos[videos["market"] == "国内"]
        fo = videos[videos["market"] != "国内"]  # 海外と言語不明
        for part in (dom, fo):
            picks.append(part.sample(n=min(per_group, len(part)), random_state=seed))
    df = pd.concat(picks).sample(frac=1, random_state=seed).reset_index(drop=True)
    df["番号"] = range(1, len(df) + 1)
    return df


def write_workbook(df, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "ラベル付け"
    headers = ["番号", "地域", "タイトル", "説明文（冒頭）", "チャンネル", "動画URL",
               "正解_観光関連か", "正解_テーマ", "正解_言語", "メモ"]
    widths = [6, 12, 60, 50, 22, 44, 14, 18, 12, 30]
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(name=FONT, bold=True)
        c.fill = YELLOW if col >= 7 else HEADER_FILL
        ws.column_dimensions[c.column_letter].width = w

    # 記入例（評価では読み飛ばす）
    example = ["例", "銚子", "【銚子電鉄】ぬれ煎餅を買いに外川駅へ", "銚子電鉄に乗って終点の外川駅まで…",
               "（例）旅チャンネル", "https://www.youtube.com/watch?v=xxxxxxxxxxx",
               "観光", "鉄道", "日本語", "電車に乗るのがメインなので鉄道"]
    for col, v in enumerate(example, start=1):
        c = ws.cell(row=2, column=col, value=v)
        c.font = Font(name=FONT, italic=True, color="808080")

    for i, r in enumerate(df.itertuples(), start=3):
        desc = (r.description if isinstance(r.description, str) else "")[:120].replace("\n", " ")
        values = [r.番号, r.region, r.title, desc, r.channel_title,
                  f"https://www.youtube.com/watch?v={r.video_id}", None, None, None, None]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=i, column=col, value=v)
            c.font = Font(name=FONT)
            c.alignment = Alignment(vertical="top", wrap_text=col in (3, 4))
            if col >= 7:
                c.fill = YELLOW
    last = len(df) + 2

    for col_letter, choices in [("G", TOURISM_CHOICES), ("H", THEME_CHOICES), ("I", LANG_CHOICES)]:
        dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        dv.error = "一覧から選んでください"
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}3:{col_letter}{last}")
    ws.freeze_panes = "C2"

    # 記入方法のシート
    guide = wb.create_sheet("記入方法")
    lines = [
        ("記入するのは黄色の4列だけです（G〜J列）。2行目は記入例なので、そのままで構いません。", True),
        ("", False),
        ("G 正解_観光関連か", True),
        ("  観光：その地域を訪れる・楽しむ内容の動画（旅行、グルメ、景色、鉄道、聖地巡礼など）", False),
        ("  観光以外：ニュース・災害、音楽、アニメやゲームの作品そのもの、学校・スポーツ、無関係な同名のものなど", False),
        ("", False),
        ("H 正解_テーマ（観光のときだけ選ぶ。観光以外なら空欄）", True),
        ("  いくつか当てはまるときは、動画の一番の主題を1つ選ぶ", False),
        ("  旅行全般：旅行やVlogだが、特定のテーマに絞れないもの", False),
        ("  どれにも当てはまらない：観光だが、一覧のどのテーマとも言えないもの", False),
        ("", False),
        ("I 正解_言語", True),
        ("  タイトルと動画の主な言語を選ぶ（字幕だけ外国語、などの場合は話している言語を優先）", False),
        ("", False),
        ("J メモ", True),
        ("  迷った理由などを自由に。迷ったら動画URLを開いて確認してよい", False),
        ("", False),
        ("※ 分類の結果はわざと載せていません。記入が終わったら evaluate_validation.py を実行してください。", True),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = guide.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT, bold=bold)
    guide.column_dimensions["A"].width = 110
    wb.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--per-group", type=int, default=15, help="地域ごと・国内/海外ごとの本数")
    parser.add_argument("--seed", type=int, default=42, help="乱数のシード")
    parser.add_argument("--region", nargs="+", default=None, help="地域（省略するとデータのある全地域）")
    parser.add_argument("--themes", nargs="+", help="このテーマに分類された動画を選ぶ（テーマ別の精度を確かめる）")
    parser.add_argument("--per-theme", type=int, default=15, help="--themes のときの、テーマごとの本数")
    parser.add_argument("--market", default="海外", choices=["海外", "国内", "すべて"], help="--themes のときの対象")
    args = parser.parse_args()

    from region_config import latest
    regions = args.region or [s for s in list_regions() if latest(s, "videos_check_*.csv") is not None]
    out_dir = DATA_ROOT / "validation"
    out_dir.mkdir(parents=True, exist_ok=True)
    if args.themes:
        df = sample_by_theme(args.themes, args.per_theme, args.market, args.seed, regions, already_sampled(out_dir))
        df["sample"] = "theme"
    else:
        df = sample_videos(args.per_group, args.seed, regions)
        df["sample"] = "random"
    stamp = datetime.now().strftime("%Y%m%d_%H%M")

    pred = df[["番号", "video_id", "region", "is_noise", "primary_theme", "themes",
               "lang_final", "market", "title", "sample"]]
    pred.to_csv(out_dir / f"predictions_{stamp}.csv", index=False, encoding="utf-8-sig")
    xlsx = out_dir / f"labeling_{stamp}.xlsx"
    write_workbook(df, xlsx)

    if args.themes:
        print(f"検証用に {len(df)}本を選びました（{args.market}向けの動画から、テーマごとに{args.per_theme}本まで。"
              f"前の検証で選んだ動画は除く）")
        print(df.groupby(["primary_theme", "region"]).size().to_string())
        print("※ どのテーマで選んだかは、記入するファイルには載せていません（先入観を避けるため）")
    else:
        print(f"検証用に {len(df)}本を選びました（地域ごとに国内・海外 各{args.per_group}本まで）")
        print(df.groupby(["region", "market"]).size().to_string())
    print(f"\nラベル付け用: {xlsx}")
    print("黄色の列に記入して上書き保存したら（または make_label_gallery.py の記入ページで記入したら）、"
          "python evaluate_validation.py --all を実行してください。")


if __name__ == "__main__":
    main()
